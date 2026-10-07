"""Build per-player appearance counts from official Project Rio game exports.

Each XLSX in games/ represents one game. Only named player rows from its Stats
sheet count; team totals and the Pitching sheet are intentionally ignored.
"""

import hashlib
from player_identity import resolve_player
import json
import sys
from collections import Counter
from pathlib import Path


ROOT = Path(__file__).resolve().parents[1]
SOURCE = Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT / "games"
OUTPUT = Path(sys.argv[2]) if len(sys.argv) > 2 else ROOT / "src" / "games-played.json"


def build_counts():
    roster = json.loads((ROOT / "src" / "players.json").read_text(encoding="utf-8"))
    ids_by_name = {player["name"].strip().casefold(): player["id"] for player in roster}
    if len(ids_by_name) != len(roster):
        raise ValueError("The player roster has duplicate names; game appearances cannot be matched safely.")

    files = sorted(SOURCE.glob("*.xlsx"))
    schedule = json.loads((ROOT / "games" / "schedule.json").read_text(encoding="utf-8"))
    overrides = {entry["report"]: entry.get("playerRows", {}) for entry in schedule if entry.get("report")}
    counts = Counter()
    fingerprints = set()
    if files:
        from openpyxl import load_workbook

    for path in files:
        fingerprint = hashlib.sha256(path.read_bytes()).hexdigest()
        if fingerprint in fingerprints:
            print(f"Skipping duplicate game file: {path.name}")
            continue
        fingerprints.add(fingerprint)
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            if "Stats" not in workbook:
                raise ValueError(f"{path.name}: missing Stats sheet")
            rows = workbook["Stats"].iter_rows(values_only=True)
            headings = next(rows, ())
            if "Player" not in headings or "Position" not in headings:
                raise ValueError(f"{path.name}: expected Player and Position columns on Stats sheet")
            player_col, position_col = headings.index("Player"), headings.index("Position")
            seen = set()
            for row_number, row in enumerate(rows, start=2):
                name = row[player_col] if len(row) > player_col else None
                position = row[position_col] if len(row) > position_col else None
                if not isinstance(name, str) or not name.strip() or not isinstance(position, str) or position.strip().upper() == "N/A":
                    continue
                scheduled_ids = next((entry["players"] for entry in schedule if entry.get("report") == path.name), None)
                player_id = overrides.get(path.name, {}).get(str(row_number)) or resolve_player(name, scheduled_ids)
                if player_id is None:
                    raise ValueError(f"{path.name}: unknown player {name!r}; correct the roster match before publishing")
                seen.add(player_id)
            if not seen:
                raise ValueError(f"{path.name}: no player appearances found")
            counts.update(seen)
            print(f"{path.name}: counted {len(seen)} players")
        finally:
            workbook.close()

    return {str(player_id): counts[player_id] for player_id in sorted(counts)}


if __name__ == "__main__":
    result = build_counts()
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    OUTPUT.write_text(json.dumps(result, indent=2) + "\n", encoding="utf-8")
    print(f"Wrote {len(result)} player game counts to {OUTPUT}")
