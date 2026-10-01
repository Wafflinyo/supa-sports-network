"""Read cached Rio XLSX or LibreOffice ODS cells without changing the report."""
from pathlib import Path
import json
from zipfile import ZipFile
from xml.etree import ElementTree as ET
from openpyxl import load_workbook

TABLE = '{urn:oasis:names:tc:opendocument:xmlns:table:1.0}'
OFFICE = '{urn:oasis:names:tc:opendocument:xmlns:office:1.0}'
TEXT = '{urn:oasis:names:tc:opendocument:xmlns:text:1.0}'


class OdsSheet:
    def __init__(self, table):
        self.table = table

    def iter_rows(self, values_only=True):
        if not values_only:
            raise ValueError('ODS reader supports cached cell values only')
        for row in self.table.iter(TABLE+'table-row'):
            values = []
            for cell in row:
                if cell.tag not in {TABLE+'table-cell', TABLE+'covered-table-cell'}:
                    continue
                kind = cell.get(OFFICE+'value-type')
                if kind in {'float', 'percentage', 'currency'}:
                    value = float(cell.get(OFFICE+'value'))
                    if value.is_integer(): value = int(value)
                elif kind == 'boolean':
                    value = cell.get(OFFICE+'boolean-value') == 'true'
                else:
                    value = '\n'.join(''.join(p.itertext()) for p in cell.findall(TEXT+'p')) or None
                count = int(cell.get(TABLE+'number-columns-repeated', '1'))
                # Rio sheets use fewer than 100 columns. Ignore LibreOffice's
                # enormous trailing empty-cell ranges, retaining internal gaps.
                values.extend([value] * min(count, max(0, 256-len(values))))
            while values and values[-1] is None: values.pop()
            if not values: continue
            count = int(row.get(TABLE+'number-rows-repeated', '1'))
            if count > 10000: raise ValueError('Unexpected repeated report rows')
            for _ in range(count): yield tuple(values)


class OdsWorkbook:
    def __init__(self, path):
        with ZipFile(path) as archive:
            root = ET.fromstring(archive.read('content.xml'))
        self.sheets = {t.get(TABLE+'name'): OdsSheet(t) for t in root.iter(TABLE+'table')}

    def __getitem__(self, name): return self.sheets[name]
    def close(self): pass


class JsonSheet:
    def __init__(self, rows): self.rows = rows
    def iter_rows(self, values_only=True):
        headings = list(self.rows[0]) if self.rows else []
        yield tuple(headings)
        for row in self.rows: yield tuple(row.get(key) for key in headings)


class JsonWorkbook:
    def __init__(self, path):
        data = json.loads(Path(path).read_text())
        self.sheets = {name: JsonSheet(data[name]) for name in ('Stats','Pitching')}
    def __getitem__(self, name): return self.sheets[name]
    def close(self): pass


def read_rio_workbook(path):
    if Path(path).suffix.lower() == '.ods': return OdsWorkbook(path)
    if Path(path).name.endswith('.rio.json'): return JsonWorkbook(path)
    return load_workbook(path, read_only=True, data_only=True)
